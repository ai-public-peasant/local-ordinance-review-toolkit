from __future__ import annotations

import argparse
import csv
import datetime as dt
import html
import json
import re
import time
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import requests
from bs4 import BeautifulSoup
from docx import Document
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


BASE_URL = "https://www.elis.go.kr"
TODAY = dt.date.today().strftime("%Y%m%d")
ARTICLE_TITLE_RE = re.compile(r"제\d+조(?:의\d+)?\([^)]+\)")
SUPPLEMENT_RE = re.compile(r"부\s*칙(?:\s*\([^)]+\)|\s*<[^>]+>)?")


@dataclass
class Entity:
    region: str
    label: str
    ctpv_cd: str
    sgg_cd: str


@dataclass
class LawItem:
    seq: int
    entity: str
    name: str
    revision_date: str
    law_type: str
    revision_type: str
    department: str
    ctpv_cd: str
    sgg_cd: str
    alr_no: str
    hist_no: str


def make_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
            )
        }
    )
    return session


def load_entities(path: Path) -> list[Entity]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    return [
        Entity(
            region=(row.get("권역") or "").strip(),
            label=(row.get("지자체명") or "").strip(),
            ctpv_cd=(row.get("ctpvCd") or "").strip(),
            sgg_cd=(row.get("sggCd") or "").strip(),
        )
        for row in rows
        if (row.get("지자체명") or "").strip()
    ]


def get_csrf(session: requests.Session, ctpv_cd: str, sgg_cd: str) -> tuple[str, str]:
    referer = f"{BASE_URL}/allalr/allAlrList?ctpvCd={ctpv_cd}&sggCd={sgg_cd}"
    response = session.get(referer, timeout=30)
    response.raise_for_status()
    match = re.search(r'name="_csrf" content="([^"]+)"', response.text)
    if not match:
        match = re.search(r'name="_csrf" value="([^"]+)"', response.text)
    if not match:
        raise RuntimeError(f"CSRF token not found: {ctpv_cd}/{sgg_cd}")
    return match.group(1), referer


def post(session: requests.Session, url: str, data: dict[str, str], csrf: str, referer: str) -> requests.Response:
    last_error: Exception | None = None
    for attempt in range(4):
        try:
            response = session.post(url, data=data, headers={"X-CSRF-TOKEN": csrf, "Referer": referer}, timeout=45)
            response.raise_for_status()
            return response
        except Exception as exc:  # noqa: BLE001
            last_error = exc
            time.sleep(1.0 * (attempt + 1))
    raise RuntimeError(f"POST failed: {url}") from last_error


def fetch_law_list(session: requests.Session, entity: Entity) -> list[LawItem]:
    csrf, referer = get_csrf(session, entity.ctpv_cd, entity.sgg_cd)
    laws: list[LawItem] = []
    seen: set[tuple[str, str]] = set()
    expected_total: int | None = None
    for page in range(1, 1000):
        data = {
            "_csrf": csrf,
            "curPage": str(page),
            "sortSe": "",
            "sortType": "",
            "srchTabSeCd": "",
            "ctpvCd": entity.ctpv_cd,
            "sggCd": entity.sgg_cd,
            "ordnRuleSe": "",
            "srchKwd": "",
            "pageSize": "100",
        }
        response = post(session, f"{BASE_URL}/allalr/allAlrList", data, csrf, referer)
        soup = BeautifulSoup(response.text, "html.parser")
        total_node = soup.select_one(".list-top .info span")
        if total_node:
            expected_total = int(total_node.get_text(strip=True).replace(",", ""))
        ids = re.findall(r"fnSrchDtls\('([^']+)',\s*'([^']+)'\)", response.text)
        rows = soup.select("div.col-table.only-pc tbody tr")
        if not rows:
            break
        for row, law_id in zip(rows, ids):
            cols = [cell.get_text(" ", strip=True) for cell in row.find_all("td")]
            if len(cols) < 6 or law_id in seen:
                continue
            seen.add(law_id)
            laws.append(
                LawItem(
                    seq=len(laws) + 1,
                    entity=cols[0] or entity.label,
                    name=cols[1].removeprefix("현 ").removeprefix("예 ").strip(),
                    revision_date=cols[2],
                    law_type=cols[3],
                    revision_type=cols[4],
                    department=cols[5],
                    ctpv_cd=entity.ctpv_cd,
                    sgg_cd=entity.sgg_cd,
                    alr_no=law_id[0],
                    hist_no=law_id[1],
                )
            )
        if expected_total is not None and len(laws) >= expected_total:
            break
        time.sleep(0.02)
    return laws


def normalize_block_text(value: str) -> str:
    value = (value or "").replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n[ \t]+", "\n", value)
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"「\s+", "「", value)
    value = re.sub(r"\s+」", "」", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def fetch_detail_articles(
    session: requests.Session,
    csrf: str,
    referer: str,
    item: LawItem,
    list_count: int,
) -> list[dict[str, str]]:
    data = {
        "_csrf": csrf,
        "curPage": "1",
        "sortSe": "",
        "srchTabSeCd": "",
        "listCnt": str(list_count),
        "ctpvCd": item.ctpv_cd,
        "sggCd": item.sgg_cd,
        "ordnRuleSe": "",
        "srchKwd": "",
        "alrNo": item.alr_no,
        "histNo": item.hist_no,
        "menuNm": "allAlr",
    }
    response = post(session, f"{BASE_URL}/allalr/selectAlrBdtOne", data, csrf, referer)
    soup = BeautifulSoup(response.text, "html.parser")
    content = soup.select_one("div.post-content")
    if content is None:
        return [{"article_id": "", "heading": "", "text": ""}]
    articles: list[dict[str, str]] = []
    for node in content.select("div.cont-area div.txt"):
        heading_node = node.select_one("strong.color-jo")
        heading = heading_node.get_text(" ", strip=True) if heading_node else ""
        text = normalize_block_text(html.unescape(node.get_text("\n", strip=True)))
        articles.append({"article_id": node.get("id", ""), "heading": heading, "text": text})
    if not articles:
        text = normalize_block_text(html.unescape(content.get_text("\n", strip=True)))
        articles.append({"article_id": "", "heading": "", "text": text})
    return articles


def clean_location(article: dict[str, str], match_start: int) -> str:
    heading = (article.get("heading") or "").strip()
    text = article.get("text") or ""
    article_id = article.get("article_id") or ""
    if heading:
        return heading
    prefix = text[: max(match_start + 1, 0)]
    article_matches = list(ARTICLE_TITLE_RE.finditer(prefix))
    if article_matches:
        title = article_matches[-1].group(0)
        return f"부칙 {title}" if article_id.startswith("BUC") else title
    supplement = SUPPLEMENT_RE.search(text[:160])
    if supplement:
        return re.sub(r"\s+", " ", supplement.group(0)).strip()
    number_title = re.search(r"(?m)^\s*(\d+)\.\s*([^\n]{1,30})", text)
    if number_title:
        return f"{number_title.group(1)}. {number_title.group(2).strip()}"[:40]
    return "부칙" if article_id.startswith("BUC") else "조문 확인 필요"


def block_text_for_output(article: dict[str, str]) -> str:
    text = normalize_block_text(article.get("text") or "")
    heading = (article.get("heading") or "").strip()
    if heading and not text.startswith(heading):
        return normalize_block_text(f"{heading}\n{text}")
    return text


def safe_cell_text(value: str) -> tuple[str, bool]:
    value = normalize_block_text(value)
    if len(value) <= 32000:
        return value, False
    trimmed = value[:31500]
    cut = max(trimmed.rfind("\n"), trimmed.rfind(". "))
    if cut > 1000:
        trimmed = trimmed[: cut + 1]
    return trimmed + "\n[원문 길이로 후략: 원문 블록이 Excel 셀 제한을 초과함]", True


def patterns_for(entity: Entity) -> list[tuple[str, str, re.Pattern[str], str, str]]:
    if entity.region == "광주":
        return [
            ("B_광역시장권한절차", "광주광역시장 권한/보고/협의", re.compile(r"광주광역시장|광주시장|(?<!구청)시장"), "상", "광역단체장 명칭 또는 권한 승계 정비 가능성"),
            ("C_광역자치법규인용", "광주광역시 자치법규 인용", re.compile(r"(?:광주광역시|광주시|시)\s*[^」\n\r]{0,90}(?:조례|규칙|훈령|예규|고시|규정)|시\s*조례|시\s*규칙"), "상", "광역 자치법규 인용사항 승계·명칭 변경 확인 필요"),
            ("E_광역기관명", "광주광역시 산하·관련 기관명", re.compile(r"광주광역시교육청|광주광역시\s*교육청|광주광역시행정심판위원회|광주광역시\s*행정심판위원회|광주광역시소청심사위원회|광주광역시\s*소청심사위원회"), "확인", "기관 존속 및 명칭 변경 여부 별도 확인 필요"),
            ("F_광역재정사업", "광주광역시비·시비·광역 보조", re.compile(r"광주광역시비|광주시비|시비\s*보조|시비|광주광역시\s*보조"), "하", "재원 명칭과 광역 지원 체계 승계 확인 필요"),
        ]
    return [
        ("B_도지사권한절차", "전라남도지사·도지사 권한/보고/협의", re.compile(r"전라남도지사|전남도지사|(?<!시)도지사"), "상", "통합특별시장 또는 신설 광역단체장 명칭으로 정비 필요 가능성"),
        ("C_도자치법규인용", "전라남도 자치법규 인용", re.compile(r"(?:전라남도|전남|도)\s*[^」\n\r]{0,90}(?:조례|규칙|훈령|예규|고시|규정)|도\s*조례|도\s*규칙"), "상", "인용 법규의 승계·명칭 변경 확인 필요"),
        ("E_도기관명", "전라남도 산하·관련 기관명", re.compile(r"전라남도교육청|전라남도\s*교육청|전라남도[가-힣]{1,12}교육지원청|전라남도행정심판위원회|전라남도\s*행정심판위원회|전라남도소청심사위원회|전라남도\s*소청심사위원회"), "확인", "기관 존속 및 명칭 변경 여부 별도 확인 필요"),
        ("F_도재정사업", "전라남도비·도비·전남 사업", re.compile(r"전라남도비|전남도비|도비\s*보조|도비|전라남도\s*보조|전라남도에서\s*추진"), "하", "재원 명칭 및 광역 지원 체계 승계 확인 필요"),
    ]


def should_skip_pattern_match(entity: Entity, code: str, term: str, block: str) -> bool:
    if code.startswith("C_"):
        # 자치구·시군 법규명 자체 또는 주소 반복은 전달용 핵심 후보에서 제외한다.
        if entity.label in term:
            return True
        if entity.label in block and term.startswith(entity.label):
            return True
    return False


def scan_law(
    entity: Entity,
    item: LawItem,
    articles: list[dict[str, str]],
    reference_law_names: list[str],
    match_reference_names: bool = False,
) -> tuple[list[dict[str, str]], int]:
    rows: list[dict[str, str]] = []
    seen: set[tuple[str, str, str, str]] = set()
    truncated_count = 0
    for article in articles:
        text = article["text"]
        block, truncated = safe_cell_text(block_text_for_output(article))
        if truncated:
            truncated_count += 1
        for code, label, pattern, priority, reason in patterns_for(entity):
            for match in pattern.finditer(text):
                term = match.group(0)
                if should_skip_pattern_match(entity, code, term, text):
                    continue
                location = clean_location(article, match.start())
                key = (item.name, location, code, block)
                if key in seen:
                    continue
                seen.add(key)
                rows.append(
                    {
                        "우선순위": priority,
                        "분류": code,
                        "검출유형": label,
                        "지자체명": entity.label,
                        "자치법규명": item.name,
                        "법규구분": item.law_type,
                        "소관부서": item.department,
                        "제개정일": item.revision_date,
                        "제개정구분": item.revision_type,
                        "조문": location,
                        "검출어": term,
                        "명칭 및 조례 인용사항": block,
                        "비고": reason,
                        "alrNo": item.alr_no,
                        "histNo": item.hist_no,
                    }
                )
        if match_reference_names:
            for ref_name in reference_law_names:
                if len(ref_name) < 8 or entity.label in ref_name:
                    continue
                idx = text.find(ref_name)
                if idx < 0:
                    continue
                location = clean_location(article, idx)
                key = (item.name, location, "G_광역법규명직접매칭", block, ref_name)
                if key in seen:
                    continue
                seen.add(key)
                rows.append(
                    {
                        "우선순위": "상",
                        "분류": "G_광역법규명직접매칭",
                        "검출유형": "광역 자치법규명 직접 매칭",
                        "지자체명": entity.label,
                        "자치법규명": item.name,
                        "법규구분": item.law_type,
                        "소관부서": item.department,
                        "제개정일": item.revision_date,
                        "제개정구분": item.revision_type,
                        "조문": location,
                        "검출어": ref_name,
                        "명칭 및 조례 인용사항": block,
                        "비고": "광역 현행 자치법규명과 직접 매칭됨",
                        "alrNo": item.alr_no,
                        "histNo": item.hist_no,
                    }
                )
    return rows, truncated_count


def style_delivery_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D9D9D9")
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 8, "B": 16, "C": 42, "D": 18, "E": 28, "F": 95, "G": 46}.items():
        sheet.column_dimensions[column].width = width
    for index in range(2, sheet.max_row + 1):
        text = str(sheet.cell(index, 6).value or "")
        lines = text.count("\n") + 1
        sheet.row_dimensions[index].height = min(max(45, lines * 18), 220)


def save_delivery_xlsx(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "제출서식"
    headers = ["연번", "지자체명", "자치법규명", "소관부서", "조문", "명칭 및 조례 인용사항", "비고"]
    sheet.append(headers)
    for index, row in enumerate(rows, start=1):
        sheet.append(
            [
                index,
                row["지자체명"],
                row["자치법규명"],
                row["소관부서"],
                row["조문"],
                row["명칭 및 조례 인용사항"],
                f"{row['검출유형']} / {row['비고']}",
            ]
        )
    style_delivery_sheet(sheet)
    workbook.save(path)


def save_entity_docx(path: Path, entity: Entity, stats: dict[str, object]) -> None:
    document = Document()
    document.add_heading(f"{entity.label} 자치법규 정비대상 제출서식 설명서", level=1)
    document.add_paragraph(f"작성일: {dt.date.today():%Y-%m-%d}")
    document.add_heading("검토 범위", level=2)
    document.add_paragraph(f"{entity.label} 현행 자치법규 상세본문을 기준으로 정비 후보를 검색했다.")
    document.add_heading("정리 기준", level=2)
    document.add_paragraph("제출용 엑셀은 작업용 시트를 제외하고 제출서식 한 시트만 작성했다.")
    document.add_paragraph("F열은 검색어 주변 발췌가 아니라 해당 조문 또는 부칙 블록 전체를 사용했고, 원문 줄바꿈을 가능한 한 보존했다.")
    document.add_paragraph("별지·별표 첨부파일 전체 텍스트 추출은 이번 1차 본문 기준 산출에 포함하지 않았다.")
    document.add_heading("검증 결과", level=2)
    for key, value in stats.items():
        document.add_paragraph(f"{key}: {value}")
    document.save(path)


def validate_rows(rows: list[dict[str, str]]) -> dict[str, int]:
    return {
        "rows": len(rows),
        "laws": len({row["자치법규명"] for row in rows}),
        "bad_location_ids": sum(1 for row in rows if re.search(r"\b(?:BUC|JOC|ART|JO)\d+", row["조문"])),
        "needs_location_check": sum(1 for row in rows if row["조문"] == "조문 확인 필요"),
        "starts_mid_sentence": sum(1 for row in rows if str(row["명칭 및 조례 인용사항"]).startswith(("시행한다.", "한다.", "따른다."))),
    }


def process_entity(
    session: requests.Session,
    project_dir: Path,
    entity: Entity,
    reference_names: dict[str, list[str]],
    force: bool,
    match_reference_names: bool,
) -> dict[str, object]:
    out_dir = project_dir / "04_제출서식" / f"전체_전달용_본문기준_{TODAY}" / entity.label
    xlsx_path = out_dir / f"{entity.label}_자치법규_정비대상_제출서식_본문기준_{TODAY}.xlsx"
    docx_path = out_dir / f"{entity.label}_자치법규_정비대상_설명서_본문기준_{TODAY}.docx"
    status_path = out_dir / f"{entity.label}_처리상태_{TODAY}.json"
    if xlsx_path.exists() and not force:
        return {"entity": entity.label, "status": "SKIPPED_EXISTS", "xlsx": str(xlsx_path)}

    laws = fetch_law_list(session, entity)
    original_law_count = len(laws)
    if entity.label == "광주광역시 서구":
        laws = [law for law in laws if law.department != "기획실"]

    csrf, referer = get_csrf(session, entity.ctpv_cd, entity.sgg_cd)
    ref_key = "광주" if entity.region == "광주" else "전남"
    candidate_rows: list[dict[str, str]] = []
    detail_failures: list[str] = []
    truncated_blocks = 0

    for index, law in enumerate(laws, start=1):
        try:
            articles = fetch_detail_articles(session, csrf, referer, law, original_law_count)
            rows, truncated = scan_law(entity, law, articles, reference_names[ref_key], match_reference_names)
            candidate_rows.extend(rows)
            truncated_blocks += truncated
        except Exception as exc:  # noqa: BLE001
            detail_failures.append(f"{law.name}: {type(exc).__name__}: {exc}")
        if index % 100 == 0:
            print(json.dumps({"entity": entity.label, "done": index, "total": len(laws), "candidates": len(candidate_rows)}, ensure_ascii=False))
        time.sleep(0.025)

    candidate_rows.sort(key=lambda row: (row["소관부서"], row["자치법규명"], row["조문"], row["명칭 및 조례 인용사항"][:80]))
    save_delivery_xlsx(xlsx_path, candidate_rows)
    validation = validate_rows(candidate_rows)
    stats = {
        "현행 자치법규 목록": original_law_count,
        "처리 대상 법규": len(laws),
        "정비 후보 행": validation["rows"],
        "후보 발생 법규": validation["laws"],
        "상세본문 실패": len(detail_failures),
        "Excel 셀 제한으로 후략한 블록": truncated_blocks,
        "조문 칸 내부 id": validation["bad_location_ids"],
        "조문 확인 필요": validation["needs_location_check"],
        "문장 중간 시작 의심": validation["starts_mid_sentence"],
    }
    save_entity_docx(docx_path, entity, stats)
    status = {"entity": entity.label, "status": "OK", "xlsx": str(xlsx_path), "docx": str(docx_path), "stats": stats, "detail_failures": detail_failures}
    status_path.write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"entity": entity.label, "status": "OK", **stats}, ensure_ascii=False))
    return status


def write_summary(project_dir: Path, results: list[dict[str, object]]) -> None:
    out_dir = project_dir / "05_총괄보고"
    out_dir.mkdir(parents=True, exist_ok=True)
    md_path = out_dir / f"전체_전달용_본문기준_처리요약_{TODAY}.md"
    csv_path = out_dir / f"전체_전달용_본문기준_처리요약_{TODAY}.csv"
    lines = [
        "# 전체 전달용 본문기준 처리요약",
        "",
        f"- 작성일: {dt.date.today():%Y-%m-%d}",
        "- 범위: 전남 22개 시군 + 광주 5개 자치구 중 완료 범위 제외",
        "- 제외: 광양시 전체 완료본, 광주광역시 서구 기획실",
        "- 기준: ELIS 상세본문 조문 기준 1차 산출, 별지·별표 첨부파일 정밀 추출 제외",
        "",
        "| 지자체 | 상태 | 후보행 | 후보법규 | 조문ID오류 | 조문확인필요 | 문장중간시작의심 |",
        "| --- | --- | ---: | ---: | ---: | ---: | ---: |",
    ]
    csv_rows: list[dict[str, object]] = []
    for result in results:
        stats = result.get("stats", {}) if isinstance(result.get("stats"), dict) else {}
        row = {
            "지자체": result.get("entity", ""),
            "상태": result.get("status", ""),
            "후보행": stats.get("정비 후보 행", ""),
            "후보법규": stats.get("후보 발생 법규", ""),
            "조문ID오류": stats.get("조문 칸 내부 id", ""),
            "조문확인필요": stats.get("조문 확인 필요", ""),
            "문장중간시작의심": stats.get("문장 중간 시작 의심", ""),
            "xlsx": result.get("xlsx", ""),
            "docx": result.get("docx", ""),
        }
        csv_rows.append(row)
        lines.append(
            f"| {row['지자체']} | {row['상태']} | {row['후보행']} | {row['후보법규']} | "
            f"{row['조문ID오류']} | {row['조문확인필요']} | {row['문장중간시작의심']} |"
        )
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(csv_rows[0].keys()) if csv_rows else ["지자체", "상태"])
        writer.writeheader()
        writer.writerows(csv_rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-dir", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--targets", type=Path, default=None)
    parser.add_argument("--force", action="store_true")
    parser.add_argument("--match-reference-names", action="store_true")
    parser.add_argument("--only", default="")
    parser.add_argument("--exclude", default="", help="Comma-separated entity labels to skip.")
    args = parser.parse_args()

    project_dir = args.project_dir
    targets = args.targets or project_dir / "data" / "targets_sample.csv"
    entities = load_entities(targets)
    only = {item.strip() for item in args.only.split(",") if item.strip()}
    if only:
        entities = [entity for entity in entities if entity.label in only]
    exclude = {item.strip() for item in args.exclude.split(",") if item.strip()}
    if exclude:
        entities = [entity for entity in entities if entity.label not in exclude]

    session = make_session()
    reference_entities = {
        "전남": Entity("전남", "전라남도", "46", "000"),
        "광주": Entity("광주", "광주광역시", "29", "000"),
    }
    reference_names = {
        key: [law.name for law in fetch_law_list(session, ref)]
        for key, ref in reference_entities.items()
    }
    print(json.dumps({"reference_laws": {key: len(value) for key, value in reference_names.items()}}, ensure_ascii=False))

    results: list[dict[str, object]] = []
    for entity in entities:
        try:
            results.append(process_entity(session, project_dir, entity, reference_names, args.force, args.match_reference_names))
        except Exception as exc:  # noqa: BLE001
            result = {"entity": entity.label, "status": "ERROR", "error": f"{type(exc).__name__}: {exc}"}
            results.append(result)
            print(json.dumps(result, ensure_ascii=False))
    write_summary(project_dir, results)


if __name__ == "__main__":
    main()
