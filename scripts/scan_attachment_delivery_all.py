from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from docx import Document
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


TODAY = dt.date.today().strftime("%Y%m%d")
SOURCE_DATE = os.environ.get("ELIS_SOURCE_DATE") or TODAY


@dataclass
class Entity:
    region: str
    label: str
    ctpv_cd: str
    sgg_cd: str


def load_entities(path: Path) -> list[Entity]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = list(csv.DictReader(file))
    return [
        Entity(
            region=(row.get("권역") or "").strip(),
            label=(row.get("지자체명") or "").strip(),
            ctpv_cd=(row.get("ctpvCd") or "").strip(),
            sgg_cd=(row.get("sggCd") or "").strip(),
        )
        for row in rows
        if (row.get("지자체명") or "").strip()
    ]


def normalize_text(value: str) -> str:
    value = (value or "").replace("\r\n", "\n").replace("\r", "\n")
    value = strip_hwp_control_garbage(value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n[ \t]+", "\n", value)
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"「\s+", "「", value)
    value = re.sub(r"\s+」", "」", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def is_ascii_packed_cjk(char: str) -> bool:
    code = ord(char)
    if not 0x4E00 <= code <= 0x9FFF:
        return False
    high = code >> 8
    low = code & 0xFF
    return 0x20 <= high <= 0x7E and 0x20 <= low <= 0x7E


def strip_hwp_control_garbage(value: str) -> str:
    result: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        packed = [char for char in buffer if is_ascii_packed_cjk(char)]
        ascii_text = "".join(chr(ord(char) >> 8) + chr(ord(char) & 0xFF) for char in packed).lower()
        if len(packed) >= 2:
            buffer.clear()
            return
        result.extend(buffer)
        buffer.clear()

    for char in value:
        if is_ascii_packed_cjk(char) or (buffer and char.isspace()):
            buffer.append(char)
        else:
            flush()
            result.append(char)
    flush()
    return "".join(result)


def safe_cell_text(value: str) -> tuple[str, bool]:
    value = normalize_text(value)
    if len(value) <= 32000:
        return value, False
    trimmed = value[:31500]
    cut = max(trimmed.rfind("\n"), trimmed.rfind(". "), trimmed.rfind("다."))
    if cut > 1000:
        trimmed = trimmed[: cut + 1]
    return trimmed + "\n[원문 길이로 후략: 별지·서식 추출텍스트가 Excel 셀 제한을 초과함]", True


def patterns_for(entity: Entity) -> list[tuple[str, str, re.Pattern[str], str, str]]:
    if entity.region == "광주":
        return [
            (
                "A_광역명칭주소",
                "광주광역시 명칭·주소 표기",
                re.compile(r"광주광역시\s*(?:동구|서구|남구|북구|광산구|청)?|광주시\s*(?:동구|서구|남구|북구|광산구)?"),
                "확인",
                "별지·서식의 광역명칭 또는 주소 표기 정비 가능성",
            ),
            (
                "B_광역시장권한절차",
                "광주광역시장 권한/보고/협의",
                re.compile(r"광주광역시장|광주시장|(?<!구청)시장"),
                "상",
                "광역단체장 명칭 또는 권한 승계 정비 가능성",
            ),
            (
                "C_광역자치법규인용",
                "광주광역시 자치법규 인용",
                re.compile(r"(?:광주광역시|광주시|시)\s*[^」\n\r]{0,90}(?:조례|규칙|훈령|예규|고시|규정)|시\s*조례|시\s*규칙"),
                "상",
                "광역 자치법규 인용사항 승계·명칭 변경 확인 필요",
            ),
            (
                "E_광역기관명",
                "광주광역시 산하·관련 기관명",
                re.compile(r"광주광역시교육청|광주광역시\s*교육청|광주광역시행정심판위원회|광주광역시\s*행정심판위원회|광주광역시소청심사위원회|광주광역시\s*소청심사위원회"),
                "확인",
                "기관 존속 및 명칭 변경 여부 별도 확인 필요",
            ),
            (
                "F_광역재정사업",
                "광주광역시비·시비·광역 보조",
                re.compile(r"광주광역시비|광주시비|시비\s*보조|시비|광주광역시\s*보조"),
                "하",
                "재원 명칭과 광역 지원 체계 승계 확인 필요",
            ),
        ]
    return [
        (
            "A_도명칭주소",
            "전라남도 명칭·주소 표기",
            re.compile(r"전라남도\s*[가-힣]{0,12}(?:시|군|청)?|전남\s*[가-힣]{0,12}(?:시|군)?"),
            "확인",
            "별지·서식의 광역명칭 또는 주소 표기 정비 가능성",
        ),
        (
            "B_도지사권한절차",
            "전라남도지사·도지사 권한/보고/협의",
            re.compile(r"전라남도지사|전남도지사|(?<!시)도지사"),
            "상",
            "통합특별시장 또는 신설 광역단체장 명칭으로 정비 필요 가능성",
        ),
        (
            "C_도자치법규인용",
            "전라남도 자치법규 인용",
            re.compile(r"(?:전라남도|전남|도)\s*[^」\n\r]{0,90}(?:조례|규칙|훈령|예규|고시|규정)|도\s*조례|도\s*규칙"),
            "상",
            "인용 법규의 승계·명칭 변경 확인 필요",
        ),
        (
            "E_도기관명",
            "전라남도 산하·관련 기관명",
            re.compile(r"전라남도교육청|전라남도\s*교육청|전라남도[가-힣]{1,12}교육지원청|전라남도행정심판위원회|전라남도\s*행정심판위원회|전라남도소청심사위원회|전라남도\s*소청심사위원회"),
            "확인",
            "기관 존속 및 명칭 변경 여부 별도 확인 필요",
        ),
        (
            "F_도재정사업",
            "전라남도비·도비·전남 사업",
            re.compile(r"전라남도비|전남도비|도비\s*보조|도비|전라남도\s*보조|전라남도에서\s*추진"),
            "하",
            "재원 명칭 및 광역 지원 체계 승계 확인 필요",
        ),
    ]


def attachment_location(label: str, text_path: str) -> str:
    label = normalize_text(label)
    if label:
        return f"별지·서식 {label}"
    return f"별지·서식 {Path(text_path).stem}"


def load_laws(entity_dir: Path, entity: Entity) -> dict[str, dict[str, str]]:
    path = entity_dir / "목록" / f"{entity.label}_현행목록_{SOURCE_DATE}.json"
    laws: dict[str, dict[str, str]] = {}
    if not path.exists():
        return laws
    for row in json.loads(path.read_text(encoding="utf-8")):
        laws[str(row.get("seq"))] = row
    return laws


def scan_entity(project_dir: Path, entity: Entity, force: bool) -> dict[str, object]:
    raw_root = project_dir / "02_원천수집" / f"ELIS_자치법규_원천자료_{SOURCE_DATE}"
    entity_dir = raw_root / entity.label
    manifest_path = entity_dir / "별지별표_첨부파일_매니페스트.csv"
    out_dir = project_dir / "04_제출서식" / f"전체_별지정밀검토_{TODAY}" / entity.label
    xlsx_path = out_dir / f"{entity.label}_자치법규_별지서식_정비후보_정밀검토_{TODAY}.xlsx"
    docx_path = out_dir / f"{entity.label}_자치법규_별지서식_정비후보_설명서_{TODAY}.docx"
    status_path = out_dir / f"{entity.label}_별지정밀검토_처리상태_{TODAY}.json"
    if xlsx_path.exists() and not force:
        return {"entity": entity.label, "status": "SKIPPED_EXISTS", "xlsx": str(xlsx_path)}
    if not manifest_path.exists():
        return {"entity": entity.label, "status": "NO_MANIFEST", "error": str(manifest_path)}

    out_dir.mkdir(parents=True, exist_ok=True)
    laws = load_laws(entity_dir, entity)
    rows: list[dict[str, str]] = []
    skipped_no_text = 0
    text_files = 0
    truncated = 0
    seen: set[tuple[str, str, str, str]] = set()

    with manifest_path.open("r", encoding="utf-8-sig", newline="") as file:
        for attachment in csv.DictReader(file):
            text_path = attachment.get("text_path") or ""
            if not text_path or not Path(text_path).exists():
                skipped_no_text += 1
                continue
            text = normalize_text(Path(text_path).read_text(encoding="utf-8", errors="ignore"))
            if not text:
                skipped_no_text += 1
                continue
            text_files += 1
            output_text, was_truncated = safe_cell_text(text)
            if was_truncated:
                truncated += 1
            law_seq = str(attachment.get("law_seq") or "")
            law = laws.get(law_seq, {})
            law_name = attachment.get("law_name") or law.get("name") or ""
            department = law.get("department") or ""
            law_type = law.get("law_type") or ""
            revision_date = law.get("revision_date") or ""
            revision_type = law.get("revision_type") or ""
            location = attachment_location(attachment.get("label") or "", text_path)

            for code, label, pattern, priority, reason in patterns_for(entity):
                matches = list(pattern.finditer(text))
                if not matches:
                    continue
                terms = sorted({normalize_text(match.group(0)) for match in matches if normalize_text(match.group(0))})
                term_text = ", ".join(terms[:8])
                if len(terms) > 8:
                    term_text += f" 외 {len(terms) - 8}개"
                key = (law_name, location, code, term_text)
                if key in seen:
                    continue
                seen.add(key)
                rows.append(
                    {
                        "우선순위": priority,
                        "분류": code,
                        "검출유형": label,
                        "지자체명": entity.label,
                        "자치법규명": law_name,
                        "법규구분": law_type,
                        "소관부서": department,
                        "제개정일": revision_date,
                        "제개정구분": revision_type,
                        "조문": location,
                        "검출어": term_text,
                        "명칭 및 조례 인용사항": output_text,
                        "비고": reason,
                        "첨부명": attachment.get("label") or "",
                        "원본파일": attachment.get("file_path") or "",
                        "추출텍스트": text_path,
                        "첨부오류": attachment.get("error") or "",
                        "alrNo": attachment.get("alr_no") or law.get("alr_no") or "",
                        "histNo": attachment.get("hist_no") or law.get("hist_no") or "",
                    }
                )

    rows.sort(key=lambda row: (row["소관부서"], row["자치법규명"], row["조문"], row["분류"], row["검출어"]))
    save_xlsx(xlsx_path, rows)
    stats = {
        "현행 자치법규 목록": len(laws),
        "검토 첨부 텍스트": text_files,
        "텍스트 없음 또는 누락": skipped_no_text,
        "별지·서식 후보 행": len(rows),
        "후보 발생 법규": len({row["자치법규명"] for row in rows}),
        "후보 발생 첨부": len({row["추출텍스트"] for row in rows}),
        "Excel 셀 제한 후략": truncated,
    }
    save_docx(docx_path, entity, stats)
    status = {"entity": entity.label, "status": "OK", "xlsx": str(xlsx_path), "docx": str(docx_path), "stats": stats}
    status_path.write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(status, ensure_ascii=False))
    return status


def save_xlsx(path: Path, rows: list[dict[str, str]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "제출서식"
    headers = ["연번", "지자체명", "자치법규명", "소관부서", "조문", "명칭 및 조례 인용사항", "비고"]
    sheet.append(headers)
    for index, row in enumerate(rows, start=1):
        sheet.append(
            [
                index,
                row["지자체명"],
                row["자치법규명"],
                row["소관부서"],
                row["조문"],
                row["명칭 및 조례 인용사항"],
                f"{row['검출유형']} / 검출어: {row['검출어']} / {row['비고']}",
            ]
        )
    style_sheet(sheet, {"A": 8, "B": 16, "C": 42, "D": 22, "E": 42, "F": 95, "G": 60})

    detail = workbook.create_sheet("별지후보상세")
    detail_headers = [
        "연번",
        "우선순위",
        "분류",
        "검출유형",
        "지자체명",
        "자치법규명",
        "법규구분",
        "소관부서",
        "조문",
        "검출어",
        "첨부명",
        "비고",
        "원본파일",
        "추출텍스트",
        "첨부오류",
        "명칭 및 조례 인용사항",
    ]
    detail.append(detail_headers)
    for index, row in enumerate(rows, start=1):
        detail.append([index] + [row.get(header, "") for header in detail_headers[1:]])
    style_sheet(
        detail,
        {
            "A": 8,
            "B": 10,
            "C": 24,
            "D": 28,
            "E": 16,
            "F": 42,
            "G": 10,
            "H": 22,
            "I": 42,
            "J": 36,
            "K": 42,
            "L": 46,
            "M": 60,
            "N": 60,
            "O": 30,
            "P": 95,
        },
    )
    workbook.save(path)


def style_sheet(sheet, widths: dict[str, int]) -> None:
    header_fill = PatternFill("solid", fgColor="E2F0D9")
    thin = Side(style="thin", color="D9D9D9")
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    text_column = "F" if sheet.title == "제출서식" else "P"
    for index in range(2, sheet.max_row + 1):
        text = str(sheet[f"{text_column}{index}"].value or "")
        lines = text.count("\n") + 1
        sheet.row_dimensions[index].height = min(max(36, lines * 16), 220)


def save_docx(path: Path, entity: Entity, stats: dict[str, object]) -> None:
    document = Document()
    document.add_heading(f"{entity.label} 별지·서식 정비후보 정밀검토 설명서", level=1)
    document.add_paragraph(f"작성일: {dt.date.today():%Y-%m-%d}")
    document.add_heading("검토 범위", level=2)
    document.add_paragraph(f"{entity.label} 자치법규 별지·별표 첨부파일 추출텍스트를 기준으로 후보를 검색했다.")
    document.add_heading("정리 기준", level=2)
    document.add_paragraph("제출서식 시트는 기존 7개 컬럼 형식을 유지하고, 별지후보상세 시트에는 첨부명·원본파일·추출텍스트 경로를 남겼다.")
    document.add_paragraph("후보는 확정 정비대상이 아니라 별지·서식 안의 광역명칭, 기관명, 재정표현, 광역 자치법규 인용 검토 대상이다.")
    document.add_heading("검증 결과", level=2)
    for key, value in stats.items():
        document.add_paragraph(f"{key}: {value}")
    document.save(path)


def write_summary(project_dir: Path, results: list[dict[str, object]]) -> None:
    out_dir = project_dir / "05_총괄보고"
    out_dir.mkdir(parents=True, exist_ok=True)
    md_path = out_dir / f"전체_별지정밀검토_처리요약_{TODAY}.md"
    csv_path = out_dir / f"전체_별지정밀검토_처리요약_{TODAY}.csv"
    lines = [
        "# 전체 별지·서식 정밀검토 처리요약",
        "",
        f"- 작성일: {dt.date.today():%Y-%m-%d}",
        f"- 기준 원천: ELIS_자치법규_원천자료_{SOURCE_DATE} 별지·별표 추출텍스트",
        "- 범위: 대상 CSV 기준",
        "",
        "| 지자체 | 상태 | 별지 후보행 | 후보법규 | 후보첨부 | 검토 첨부텍스트 | 텍스트 누락 | 후략 |",
        "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: |",
    ]
    csv_rows: list[dict[str, object]] = []
    for result in results:
        stats = result.get("stats", {}) if isinstance(result.get("stats"), dict) else {}
        row = {
            "지자체": result.get("entity", ""),
            "상태": result.get("status", ""),
            "별지후보행": stats.get("별지·서식 후보 행", ""),
            "후보법규": stats.get("후보 발생 법규", ""),
            "후보첨부": stats.get("후보 발생 첨부", ""),
            "검토첨부텍스트": stats.get("검토 첨부 텍스트", ""),
            "텍스트누락": stats.get("텍스트 없음 또는 누락", ""),
            "후략": stats.get("Excel 셀 제한 후략", ""),
            "xlsx": result.get("xlsx", ""),
            "docx": result.get("docx", ""),
            "error": result.get("error", ""),
        }
        csv_rows.append(row)
        lines.append(
            f"| {row['지자체']} | {row['상태']} | {row['별지후보행']} | {row['후보법규']} | "
            f"{row['후보첨부']} | {row['검토첨부텍스트']} | {row['텍스트누락']} | {row['후략']} |"
        )
    md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(csv_rows[0].keys()) if csv_rows else ["지자체", "상태"])
        writer.writeheader()
        writer.writerows(csv_rows)


def main() -> None:
    global SOURCE_DATE
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-dir", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--targets", type=Path, default=None)
    parser.add_argument("--only", default="")
    parser.add_argument("--exclude", default="", help="Comma-separated entity labels to skip.")
    parser.add_argument("--source-date", default=SOURCE_DATE, help="Raw collection date folder suffix, e.g. 20260511.")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    SOURCE_DATE = args.source_date

    project_dir = args.project_dir
    targets = args.targets or project_dir / "data" / "targets_sample.csv"
    entities = load_entities(targets)
    only = {item.strip() for item in args.only.split(",") if item.strip()}
    if only:
        entities = [entity for entity in entities if entity.label in only]
    exclude = {item.strip() for item in args.exclude.split(",") if item.strip()}
    if exclude:
        entities = [entity for entity in entities if entity.label not in exclude]

    results: list[dict[str, object]] = []
    for entity in entities:
        try:
            results.append(scan_entity(project_dir, entity, args.force))
        except Exception as exc:  # noqa: BLE001
            result = {"entity": entity.label, "status": "ERROR", "error": f"{type(exc).__name__}: {exc}"}
            results.append(result)
            print(json.dumps(result, ensure_ascii=False))
    write_summary(project_dir, results)


if __name__ == "__main__":
    main()
