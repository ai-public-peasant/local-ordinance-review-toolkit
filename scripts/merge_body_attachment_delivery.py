from __future__ import annotations

import csv
import datetime as dt
import json
import re
import argparse
import os
from pathlib import Path

import openpyxl
from docx import Document
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


TODAY = dt.date.today().strftime("%Y%m%d")
BODY_DATE = os.environ.get("ELIS_BODY_DATE") or TODAY
ATTACHMENT_DATE = os.environ.get("ELIS_ATTACHMENT_DATE") or TODAY


def is_ascii_packed_cjk(char: str) -> bool:
    code = ord(char)
    if not 0x4E00 <= code <= 0x9FFF:
        return False
    high = code >> 8
    low = code & 0xFF
    return 0x20 <= high <= 0x7E and 0x20 <= low <= 0x7E


def strip_hwp_control_garbage(value: str) -> str:
    result: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        packed = [char for char in buffer if is_ascii_packed_cjk(char)]
        if len(packed) >= 2:
            buffer.clear()
            return
        result.extend(buffer)
        buffer.clear()

    for char in value or "":
        if is_ascii_packed_cjk(char) or (buffer and char.isspace()):
            buffer.append(char)
        else:
            flush()
            result.append(char)
    flush()
    return "".join(result)


def clean_cell(value: object) -> str:
    text = "" if value is None else str(value)
    text = strip_hwp_control_garbage(text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def read_delivery_rows(path: Path, source: str) -> list[dict[str, str]]:
    if not path.exists():
        return []
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["제출서식"]
    rows: list[dict[str, str]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        rows.append(
            {
                "구분": source,
                "지자체명": clean_cell(row[1]),
                "자치법규명": clean_cell(row[2]),
                "소관부서": clean_cell(row[3]),
                "조문": clean_cell(row[4]),
                "명칭 및 조례 인용사항": clean_cell(row[5]),
                "비고": f"{source} / {clean_cell(row[6])}",
            }
        )
    workbook.close()
    return rows


def style_sheet(sheet, widths: dict[str, int], text_column: str = "F") -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D9D9D9")
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for index in range(2, sheet.max_row + 1):
        text = str(sheet[f"{text_column}{index}"].value or "")
        sheet.row_dimensions[index].height = min(max(36, (text.count("\n") + 1) * 15), 220)


def save_entity_xlsx(path: Path, rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "제출서식"
    headers = ["연번", "지자체명", "자치법규명", "소관부서", "조문", "명칭 및 조례 인용사항", "비고"]
    sheet.append(headers)
    for index, row in enumerate(rows, start=1):
        sheet.append(
            [
                index,
                row["지자체명"],
                row["자치법규명"],
                row["소관부서"],
                row["조문"],
                row["명칭 및 조례 인용사항"],
                row["비고"],
            ]
        )
    style_sheet(sheet, {"A": 8, "B": 16, "C": 42, "D": 22, "E": 42, "F": 95, "G": 62})

    detail = workbook.create_sheet("통합상세")
    detail_headers = ["연번", "구분", "지자체명", "자치법규명", "소관부서", "조문", "비고", "명칭 및 조례 인용사항"]
    detail.append(detail_headers)
    for index, row in enumerate(rows, start=1):
        detail.append([index, row["구분"], row["지자체명"], row["자치법규명"], row["소관부서"], row["조문"], row["비고"], row["명칭 및 조례 인용사항"]])
    style_sheet(detail, {"A": 8, "B": 12, "C": 16, "D": 42, "E": 22, "F": 42, "G": 62, "H": 95}, "H")
    workbook.save(path)


def save_entity_docx(path: Path, entity: str, stats: dict[str, int]) -> None:
    document = Document()
    document.add_heading(f"{entity} 자치법규 본문+별지 통합 정비후보 설명서", level=1)
    document.add_paragraph(f"작성일: {dt.date.today():%Y-%m-%d}")
    document.add_heading("검토 범위", level=2)
    document.add_paragraph("1차 본문기준 제출서식과 별지·서식 정밀검토 제출서식을 하나의 제출서식으로 합쳤다.")
    document.add_heading("검증 결과", level=2)
    for key, value in stats.items():
        document.add_paragraph(f"{key}: {value}")
    document.save(path)


def main() -> None:
    global BODY_DATE, ATTACHMENT_DATE
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-dir", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--targets", type=Path, default=None)
    parser.add_argument("--body-date", default=BODY_DATE)
    parser.add_argument("--attachment-date", default=ATTACHMENT_DATE)
    parser.add_argument("--only", default="")
    parser.add_argument("--exclude", default="", help="Comma-separated entity labels to skip.")
    args = parser.parse_args()

    project_dir = args.project_dir
    BODY_DATE = args.body_date
    ATTACHMENT_DATE = args.attachment_date
    targets = args.targets or project_dir / "data" / "targets_sample.csv"
    with targets.open("r", encoding="utf-8-sig", newline="") as file:
        entities = [(row.get("지자체명") or "").strip() for row in csv.DictReader(file)]
    entities = [entity for entity in entities if entity]
    only = {item.strip() for item in args.only.split(",") if item.strip()}
    exclude = {item.strip() for item in args.exclude.split(",") if item.strip()}
    if only:
        entities = [entity for entity in entities if entity in only]
    if exclude:
        entities = [entity for entity in entities if entity not in exclude]

    body_root = project_dir / "04_제출서식" / f"전체_전달용_본문기준_{BODY_DATE}"
    attachment_root = project_dir / "04_제출서식" / f"전체_별지정밀검토_{ATTACHMENT_DATE}"
    out_root = project_dir / "04_제출서식" / f"전체_통합검토_본문별지_{TODAY}"
    summary_dir = project_dir / "05_총괄보고"
    summary_dir.mkdir(parents=True, exist_ok=True)

    all_rows: list[dict[str, str]] = []
    summaries: list[dict[str, object]] = []
    for entity in entities:
        body_path = body_root / entity / f"{entity}_자치법규_정비대상_제출서식_본문기준_{BODY_DATE}.xlsx"
        attachment_path = attachment_root / entity / f"{entity}_자치법규_별지서식_정비후보_정밀검토_{ATTACHMENT_DATE}.xlsx"
        body_rows = read_delivery_rows(body_path, "본문")
        attachment_rows = read_delivery_rows(attachment_path, "별지·서식")
        rows = body_rows + attachment_rows
        rows.sort(key=lambda row: (row["소관부서"], row["자치법규명"], row["구분"], row["조문"], row["명칭 및 조례 인용사항"][:80]))

        entity_dir = out_root / entity
        xlsx_path = entity_dir / f"{entity}_자치법규_정비후보_본문별지_통합_{TODAY}.xlsx"
        docx_path = entity_dir / f"{entity}_자치법규_정비후보_본문별지_통합_설명서_{TODAY}.docx"
        save_entity_xlsx(xlsx_path, rows)
        stats = {
            "본문 후보 행": len(body_rows),
            "별지·서식 후보 행": len(attachment_rows),
            "통합 후보 행": len(rows),
            "후보 발생 법규": len({row["자치법규명"] for row in rows}),
            "본문 원본 파일 존재": int(body_path.exists()),
            "별지 원본 파일 존재": int(attachment_path.exists()),
        }
        save_entity_docx(docx_path, entity, stats)
        summaries.append({"지자체": entity, **stats, "xlsx": str(xlsx_path), "docx": str(docx_path)})
        all_rows.extend(rows)

    master_path = summary_dir / f"전체_본문별지_통합마스터_{TODAY}.xlsx"
    save_entity_xlsx(master_path, all_rows)
    summary_csv = summary_dir / f"전체_본문별지_통합검토_처리요약_{TODAY}.csv"
    with summary_csv.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(summaries[0].keys()))
        writer.writeheader()
        writer.writerows(summaries)

    body_total = sum(int(row["본문 후보 행"]) for row in summaries)
    attachment_total = sum(int(row["별지·서식 후보 행"]) for row in summaries)
    total = sum(int(row["통합 후보 행"]) for row in summaries)
    md_lines = [
        "# 전체 본문+별지 통합검토 처리요약",
        "",
        f"- 작성일: {dt.date.today():%Y-%m-%d}",
        "- 범위: 대상 CSV 기준",
        f"- 처리 대상: {len(summaries)}개 지자체",
        f"- 본문 후보 행: {body_total:,}건",
        f"- 별지·서식 후보 행: {attachment_total:,}건",
        f"- 통합 후보 행: {total:,}건",
        f"- 통합 마스터: `05_총괄보고\\{master_path.name}`",
        f"- 지자체별 산출 폴더: `04_제출서식\\{out_root.name}`",
        "",
        "| 지자체 | 본문 후보행 | 별지·서식 후보행 | 통합 후보행 | 후보 법규 |",
        "| --- | ---: | ---: | ---: | ---: |",
    ]
    for row in summaries:
        md_lines.append(
            f"| {row['지자체']} | {row['본문 후보 행']} | {row['별지·서식 후보 행']} | {row['통합 후보 행']} | {row['후보 발생 법규']} |"
        )
    summary_md = summary_dir / f"전체_본문별지_통합검토_처리요약_{TODAY}.md"
    summary_md.write_text("\n".join(md_lines) + "\n", encoding="utf-8")

    print(json.dumps({"entities": len(summaries), "body_rows": body_total, "attachment_rows": attachment_total, "total_rows": total, "master": str(master_path), "folder": str(out_root)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
