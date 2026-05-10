from __future__ import annotations

import argparse
import csv
import datetime as dt
import html
import json
import os
import re
import shutil
import struct
import time
import unicodedata
import zipfile
import zlib
from dataclasses import asdict, dataclass
from pathlib import Path
from urllib.parse import unquote

import olefile
import openpyxl
import requests
from bs4 import BeautifulSoup
from openpyxl.styles import Alignment, Font, PatternFill


BASE_URL = "https://www.elis.go.kr"
TODAY = os.environ.get("ELIS_COLLECT_DATE") or dt.date.today().strftime("%Y%m%d")


@dataclass
class Entity:
    label: str
    ctpv_cd: str
    sgg_cd: str


@dataclass
class LawItem:
    seq: int
    entity: str
    name: str
    revision_date: str
    law_type: str
    revision_type: str
    department: str
    ctpv_cd: str
    sgg_cd: str
    alr_no: str
    hist_no: str


@dataclass
class AttachmentItem:
    entity: str
    law_seq: int
    law_name: str
    alr_no: str
    hist_no: str
    attlist_sn: str
    label: str
    downloaded: bool = False
    file_path: str = ""
    text_path: str = ""
    byte_size: int = 0
    text_chars: int = 0
    error: str = ""


ENTITIES = [
    Entity("광양시", "46", "230"),
    Entity("전라남도", "46", "000"),
]


def load_entities_from_csv(path: Path, only: set[str] | None = None) -> list[Entity]:
    entities: list[Entity] = []
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        for row in csv.DictReader(file):
            label = (row.get("지자체명") or "").strip()
            if not label:
                continue
            if only and label not in only:
                continue
            entities.append(
                Entity(
                    label=label,
                    ctpv_cd=(row.get("ctpvCd") or "").strip(),
                    sgg_cd=(row.get("sggCd") or "").strip(),
                )
            )
    if not entities:
        raise ValueError(f"No entities loaded from {path}")
    return entities


def safe_filename(value: str, limit: int = 120) -> str:
    value = unicodedata.normalize("NFKC", value or "")
    value = re.sub(r'[<>:"/\\|?*\n\r\t]+', "_", value).strip()
    value = re.sub(r"\s+", " ", value)
    return value[:limit].rstrip(" .") or "untitled"


def is_ascii_packed_cjk(char: str) -> bool:
    code = ord(char)
    if not 0x4E00 <= code <= 0x9FFF:
        return False
    high = code >> 8
    low = code & 0xFF
    return 0x20 <= high <= 0x7E and 0x20 <= low <= 0x7E


def strip_hwp_control_garbage(value: str) -> str:
    result: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        packed = [char for char in buffer if is_ascii_packed_cjk(char)]
        ascii_text = "".join(chr(ord(char) >> 8) + chr(ord(char) & 0xFF) for char in packed).lower()
        if len(packed) >= 2:
            buffer.clear()
            return
        result.extend(buffer)
        buffer.clear()

    for char in value:
        if is_ascii_packed_cjk(char) or (buffer and char.isspace()):
            buffer.append(char)
        else:
            flush()
            result.append(char)
    flush()
    return "".join(result)


def clean_text(value: str) -> str:
    value = html.unescape(value or "")
    value = value.encode("utf-8", errors="ignore").decode("utf-8")
    value = strip_hwp_control_garbage(value)
    value = value.replace("\u0000", " ")
    value = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f]+", " ", value)
    value = re.sub(r"[ \t\r\f\v]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def make_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"
            )
        }
    )
    return session


def get_csrf(session: requests.Session, entity: Entity) -> tuple[str, str]:
    referer = f"{BASE_URL}/allalr/allAlrList?ctpvCd={entity.ctpv_cd}&sggCd={entity.sgg_cd}"
    response = session.get(referer, timeout=45)
    response.raise_for_status()
    match = re.search(r'name="_csrf" content="([^"]+)"', response.text)
    if not match:
        match = re.search(r'name="_csrf" value="([^"]+)"', response.text)
    if not match:
        raise RuntimeError(f"CSRF not found for {entity.label}")
    return match.group(1), referer


def post_with_retry(
    session: requests.Session,
    url: str,
    data: dict[str, str],
    csrf: str,
    referer: str,
    retries: int = 5,
) -> requests.Response:
    last_error: Exception | None = None
    for attempt in range(retries):
        try:
            response = session.post(
                url,
                data=data,
                headers={"X-CSRF-TOKEN": csrf, "Referer": referer},
                timeout=60,
            )
            response.raise_for_status()
            return response
        except Exception as exc:  # noqa: BLE001 - ELIS often has transient resets.
            last_error = exc
            time.sleep(1.2 * (attempt + 1))
    raise RuntimeError(f"POST failed: {url}") from last_error


def fetch_law_list(session: requests.Session, csrf: str, referer: str, entity: Entity) -> list[LawItem]:
    laws: list[LawItem] = []
    seen: set[tuple[str, str]] = set()
    expected_total: int | None = None

    for page in range(1, 1000):
        data = {
            "_csrf": csrf,
            "curPage": str(page),
            "sortSe": "",
            "sortType": "",
            "srchTabSeCd": "",
            "ctpvCd": entity.ctpv_cd,
            "sggCd": entity.sgg_cd,
            "ordnRuleSe": "",
            "srchKwd": "",
            "pageSize": "100",
        }
        response = post_with_retry(session, f"{BASE_URL}/allalr/allAlrList", data, csrf, referer)
        soup = BeautifulSoup(response.text, "html.parser")
        total_node = soup.select_one(".list-top .info span")
        if total_node:
            total_text = re.sub(r"\D+", "", total_node.get_text(strip=True))
            if total_text:
                expected_total = int(total_text)

        ids = re.findall(r"fnSrchDtls\('([^']+)',\s*'([^']+)'\)", response.text)
        rows = soup.select("div.col-table.only-pc tbody tr")
        if not rows:
            break

        before_count = len(laws)
        for row, law_id in zip(rows, ids):
            cols = [cell.get_text(" ", strip=True) for cell in row.find_all("td")]
            if len(cols) < 6:
                continue
            if law_id in seen:
                continue
            seen.add(law_id)
            laws.append(
                LawItem(
                    seq=len(laws) + 1,
                    entity=cols[0] or entity.label,
                    name=cols[1].removeprefix("현 ").strip(),
                    revision_date=cols[2],
                    law_type=cols[3],
                    revision_type=cols[4],
                    department=cols[5],
                    ctpv_cd=entity.ctpv_cd,
                    sgg_cd=entity.sgg_cd,
                    alr_no=law_id[0],
                    hist_no=law_id[1],
                )
            )

        if len(laws) == before_count:
            break
        if expected_total is not None and len(laws) >= expected_total:
            break
        time.sleep(0.05)

    return laws


def fetch_detail_html(
    session: requests.Session,
    csrf: str,
    referer: str,
    law: LawItem,
    list_count: int,
) -> str:
    data = {
        "_csrf": csrf,
        "curPage": "1",
        "sortSe": "",
        "srchTabSeCd": "",
        "listCnt": str(list_count),
        "ctpvCd": law.ctpv_cd,
        "sggCd": law.sgg_cd,
        "ordnRuleSe": "",
        "srchKwd": "",
        "alrNo": law.alr_no,
        "histNo": law.hist_no,
        "menuNm": "allAlr",
    }
    response = post_with_retry(session, f"{BASE_URL}/allalr/selectAlrBdtOne", data, csrf, referer)
    return response.text


def extract_body_text(detail_html: str) -> str:
    soup = BeautifulSoup(detail_html, "html.parser")
    content = soup.select_one("div.post-content") or soup
    return clean_text(content.get_text("\n", strip=True))


def parse_attachments(detail_html: str, law: LawItem) -> list[AttachmentItem]:
    soup = BeautifulSoup(detail_html, "html.parser")
    attachments: list[AttachmentItem] = []
    seen: set[str] = set()
    for link in soup.select('a[href*="fnAttListDown"], a[onclick*="fnAttListDown"]'):
        href = " ".join([link.get("href") or "", link.get("onclick") or ""])
        match = re.search(r"fnAttListDown\('([^']+)'\s*,\s*'([^']+)'\s*,\s*'([^']+)'\)", href)
        if not match:
            continue
        alr_no, hist_no, attlist_sn = match.groups()
        if attlist_sn in seen:
            continue
        seen.add(attlist_sn)
        label = clean_text(link.get_text(" ", strip=True))
        attachments.append(
            AttachmentItem(
                entity=law.entity,
                law_seq=law.seq,
                law_name=law.name,
                alr_no=alr_no,
                hist_no=hist_no,
                attlist_sn=attlist_sn,
                label=label or f"첨부_{attlist_sn}",
            )
        )
    return attachments


def decode_content_disposition_filename(header_value: str | None) -> str:
    if not header_value:
        return ""
    star = re.search(r"filename\*=([^']*)''([^;]+)", header_value, flags=re.I)
    if star:
        return unquote(star.group(2))
    match = re.search(r'filename="([^"]+)"', header_value, flags=re.I)
    if not match:
        match = re.search(r"filename=([^;]+)", header_value, flags=re.I)
    if not match:
        return ""
    name = match.group(1).strip().strip('"')
    try:
        name = name.encode("latin-1").decode("utf-8")
    except Exception:  # noqa: BLE001 - keep original if the server already decoded it.
        pass
    return name


def extension_from_content(content: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix:
        return suffix
    if content.startswith(b"\xd0\xcf\x11\xe0"):
        return ".hwp"
    if content.startswith(b"PK\x03\x04"):
        return ".zip"
    if content.startswith(b"%PDF"):
        return ".pdf"
    return ".bin"


def download_attachment(
    session: requests.Session,
    csrf: str,
    referer: str,
    attachment: AttachmentItem,
    file_dir: Path,
) -> AttachmentItem:
    file_dir.mkdir(parents=True, exist_ok=True)
    ajax = post_with_retry(
        session,
        f"{BASE_URL}/allalr/attListDownAajx",
        {
            "alrNo": attachment.alr_no,
            "histNo": attachment.hist_no,
            "attlistSn": attachment.attlist_sn,
        },
        csrf,
        referer,
    )
    if ajax.text.strip().lower() not in {"true", "1", "y", "yes"}:
        raise RuntimeError(f"attachment ajax denied: {ajax.text[:100]}")

    chk_time = dt.datetime.now().strftime("%Y%m%d%H%M%S%f")[:17]
    url = (
        f"{BASE_URL}/allalr/attListDown?alrNo={attachment.alr_no}"
        f"&histNo={attachment.hist_no}&attlistSn={attachment.attlist_sn}&chkTime={chk_time}"
    )
    response = session.get(url, headers={"Referer": referer}, timeout=90)
    response.raise_for_status()
    filename = decode_content_disposition_filename(response.headers.get("content-disposition"))
    ext = extension_from_content(response.content, filename)
    base_name = safe_filename(f"{attachment.attlist_sn}_{attachment.label}_{filename or ''}", 150)
    if not base_name.lower().endswith(ext):
        base_name += ext
    file_path = file_dir / base_name
    file_path.write_bytes(response.content)
    attachment.downloaded = True
    attachment.file_path = str(file_path)
    attachment.byte_size = len(response.content)
    return attachment


def extract_hwp_text(path: Path) -> str:
    ole = olefile.OleFileIO(str(path))
    header = ole.openstream("FileHeader").read()
    compressed = bool(header[36] & 1)
    texts: list[str] = []
    for stream_path in ole.listdir():
        if len(stream_path) != 2 or stream_path[0] != "BodyText" or not stream_path[1].startswith("Section"):
            continue
        data = ole.openstream("/".join(stream_path)).read()
        if compressed:
            data = zlib.decompress(data, -15)
        pos = 0
        while pos + 4 <= len(data):
            header_value = struct.unpack_from("<I", data, pos)[0]
            pos += 4
            tag_id = header_value & 0x3FF
            size = (header_value >> 20) & 0xFFF
            if size == 0xFFF:
                if pos + 4 > len(data):
                    break
                size = struct.unpack_from("<I", data, pos)[0]
                pos += 4
            record = data[pos : pos + size]
            pos += size
            if tag_id == 67:
                texts.append(record.decode("utf-16le", errors="ignore"))
    return clean_text("\n".join(texts))


def extract_zip_text(path: Path) -> str:
    parts: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            lower = name.lower()
            if lower.endswith((".xml", ".txt", ".rels")):
                try:
                    raw = archive.read(name)
                except Exception:
                    continue
                text = raw.decode("utf-8", errors="ignore")
                text = re.sub(r"<[^>]+>", " ", text)
                parts.append(text)
    return clean_text("\n".join(parts))


def extract_pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        return ""
    reader = PdfReader(str(path))
    return clean_text("\n".join(page.extract_text() or "" for page in reader.pages))


def extract_office_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".hwp" or path.read_bytes()[:4] == b"\xd0\xcf\x11\xe0":
        return extract_hwp_text(path)
    if suffix in {".hwpx", ".docx", ".xlsx", ".zip"} or path.read_bytes()[:4] == b"PK\x03\x04":
        return extract_zip_text(path)
    if suffix == ".pdf" or path.read_bytes()[:4] == b"%PDF":
        return extract_pdf_text(path)
    try:
        return clean_text(path.read_text(encoding="utf-8", errors="ignore"))
    except Exception:
        return ""


def write_law_list(work_dir: Path, entity: Entity, laws: list[LawItem]) -> None:
    list_dir = work_dir / entity.label / "목록"
    list_dir.mkdir(parents=True, exist_ok=True)
    (list_dir / f"{entity.label}_현행목록_{TODAY}.json").write_text(
        json.dumps([asdict(law) for law in laws], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "현행목록"
    headers = ["연번", "자치단체", "자치법규명", "제개정일", "법규구분", "제개정구분", "소관부서", "alrNo", "histNo"]
    sheet.append(headers)
    for law in laws:
        sheet.append(
            [
                law.seq,
                law.entity,
                law.name,
                law.revision_date,
                law.law_type,
                law.revision_type,
                law.department,
                law.alr_no,
                law.hist_no,
            ]
        )
    style_sheet(sheet)
    workbook.save(list_dir / f"{entity.label}_현행목록_{TODAY}.xlsx")


def style_sheet(sheet: openpyxl.worksheet.worksheet.Worksheet) -> None:
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = 10
        for cell in column[:300]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            width = max(width, min(len(str(cell.value or "")) + 2, 70))
        sheet.column_dimensions[letter].width = width


def write_manifest(path: Path, rows: list[AttachmentItem]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(asdict(rows[0]).keys()) if rows else list(AttachmentItem("", 0, "", "", "", "", "").__dict__.keys())
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(asdict(row))


def write_summary(path: Path, summary: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = ["# 자치법규 전수 수집 검증 요약", ""]
    for key, value in summary.items():
        lines.append(f"- {key}: {value}")
    path.write_text("\n".join(lines), encoding="utf-8")


def collect_entity(project_dir: Path, entity: Entity) -> dict[str, object]:
    work_dir = project_dir / "02_원천수집" / f"ELIS_자치법규_원천자료_{TODAY}"
    entity_dir = work_dir / entity.label
    done_path = entity_dir / f"_DONE_{TODAY}.json"
    if done_path.exists():
        return json.loads(done_path.read_text(encoding="utf-8"))
    html_dir = entity_dir / "상세HTML"
    body_dir = entity_dir / "본문텍스트"
    att_dir = entity_dir / "별지별표_원본파일"
    att_text_dir = entity_dir / "별지별표_추출텍스트"
    merged_dir = entity_dir / "통합검색텍스트"
    for folder in [html_dir, body_dir, att_dir, att_text_dir, merged_dir]:
        folder.mkdir(parents=True, exist_ok=True)

    session = make_session()
    csrf, referer = get_csrf(session, entity)
    laws = fetch_law_list(session, csrf, referer, entity)
    write_law_list(work_dir, entity, laws)

    attachments: list[AttachmentItem] = []
    detail_failures: list[str] = []
    for law in laws:
        prefix = f"{law.seq:04d}_{safe_filename(law.name, 110)}"
        html_path = html_dir / f"{prefix}.html"
        body_path = body_dir / f"{prefix}.txt"
        if html_path.exists():
            detail_html = html_path.read_text(encoding="utf-8")
        else:
            try:
                detail_html = fetch_detail_html(session, csrf, referer, law, len(laws))
                html_path.write_text(detail_html, encoding="utf-8")
                time.sleep(0.04)
            except Exception as exc:  # noqa: BLE001
                detail_failures.append(f"{law.seq} {law.name}: {exc}")
                continue
        body_text = extract_body_text(detail_html)
        body_path.write_text(body_text, encoding="utf-8")

        law_attachments = parse_attachments(detail_html, law)
        law_att_dir = att_dir / prefix
        law_att_text_dir = att_text_dir / prefix
        law_att_text_dir.mkdir(parents=True, exist_ok=True)
        for attachment in law_attachments:
            text_path = law_att_text_dir / f"{attachment.attlist_sn}_{safe_filename(attachment.label, 80)}.txt"
            try:
                if not attachment.file_path:
                    existing = sorted(law_att_dir.glob(f"{attachment.attlist_sn}_*"))
                    if existing:
                        attachment.downloaded = True
                        attachment.file_path = str(existing[0])
                        attachment.byte_size = existing[0].stat().st_size
                    else:
                        attachment = download_attachment(session, csrf, referer, attachment, law_att_dir)
                        time.sleep(0.04)
                source_path = Path(attachment.file_path)
                if text_path.exists():
                    text = text_path.read_text(encoding="utf-8")
                else:
                    text = extract_office_text(source_path)
                    text_path.write_text(text, encoding="utf-8")
                attachment.text_path = str(text_path)
                attachment.text_chars = len(text)
            except Exception as exc:  # noqa: BLE001
                attachment.error = str(exc)
            attachments.append(attachment)

        merged_parts = [body_text]
        for attachment in law_attachments:
            if attachment.text_path and Path(attachment.text_path).exists():
                merged_parts.append(Path(attachment.text_path).read_text(encoding="utf-8"))
        (merged_dir / f"{prefix}.txt").write_text(clean_text("\n\n".join(merged_parts)), encoding="utf-8")

        if law.seq % 100 == 0:
            print(
                json.dumps(
                    {
                        "entity": entity.label,
                        "laws_done": law.seq,
                        "laws_total": len(laws),
                        "attachments": len(attachments),
                        "attachment_errors": sum(1 for item in attachments if item.error),
                    },
                    ensure_ascii=False,
                )
            )

    manifest_path = entity_dir / "별지별표_첨부파일_매니페스트.csv"
    write_manifest(manifest_path, attachments)
    summary = {
        "entity": entity.label,
        "laws": len(laws),
        "attachments_found": len(attachments),
        "attachments_downloaded": sum(1 for item in attachments if item.downloaded),
        "attachments_text_extracted": sum(1 for item in attachments if item.text_chars > 0),
        "attachment_errors": sum(1 for item in attachments if item.error),
        "detail_failures": len(detail_failures),
        "manifest": str(manifest_path),
    }
    done_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def copy_input_documents(project_dir: Path) -> None:
    target = project_dir / "00_공문_제출서식_원본"
    target.mkdir(exist_ok=True)
    for pattern in ["*.pdf", "*.xlsx", "*.png"]:
        for source in project_dir.glob(pattern):
            if source.parent == target:
                continue
            destination = target / source.name
            if not destination.exists():
                shutil.copy2(source, destination)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-dir", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--entities-csv", type=Path, default=None)
    parser.add_argument("--force", action="store_true", help="Recollect even if _DONE marker exists.")
    parser.add_argument(
        "--only",
        default="",
        help="Comma-separated entity labels to collect, e.g. 여수시,해남군,광주광역시 남구",
    )
    args = parser.parse_args()

    project_dir = args.project_dir
    entities_csv = args.entities_csv or project_dir / "data" / "targets_sample.csv"
    only = {item.strip() for item in args.only.split(",") if item.strip()} or None
    entities = load_entities_from_csv(entities_csv, only=only) if entities_csv.exists() else ENTITIES

    if args.force:
        work_dir = project_dir / "02_원천수집" / f"ELIS_자치법규_원천자료_{TODAY}"
        for entity in entities:
            done_path = work_dir / entity.label / f"_DONE_{TODAY}.json"
            if done_path.exists():
                done_path.unlink()

    copy_input_documents(project_dir)
    summaries = [collect_entity(project_dir, entity) for entity in entities]
    total_summary = {
        "검증일": TODAY,
        "수집대상": ", ".join(entity.label for entity in entities),
        "총 자치법규 수": sum(int(item["laws"]) for item in summaries),
        "총 별지별표 첨부 발견": sum(int(item["attachments_found"]) for item in summaries),
        "총 별지별표 첨부 다운로드": sum(int(item["attachments_downloaded"]) for item in summaries),
        "총 별지별표 텍스트 추출": sum(int(item["attachments_text_extracted"]) for item in summaries),
        "총 첨부 오류": sum(int(item["attachment_errors"]) for item in summaries),
        "상세 실패": sum(int(item["detail_failures"]) for item in summaries),
        "개별요약": json.dumps(summaries, ensure_ascii=False),
    }
    verify_dir = project_dir / "05_총괄보고"
    write_summary(verify_dir / f"전수수집_검증요약_{TODAY}.md", total_summary)
    print(json.dumps(total_summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
